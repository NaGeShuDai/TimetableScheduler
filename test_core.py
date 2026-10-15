import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from core import Person, SLOTS, parse_timetable, parse_course_cell, weeks_from_text, available, candidates, solve_schedule, export_grid, compact_weeks, total_free_cells


def people():
    return [Person(str(i), name, '2026-2027-1', '', '', notes_confirmed=True)
            for i, name in enumerate(['甲', '乙', '丙'])]


def rule(kind, pid, day=0, slot=0, weeks=None):
    return dict(kind=kind, person=pid, day=day, slot=slot, weeks=weeks or [], enabled=True)


class CoreTests(unittest.TestCase):
    def test_week_word_in_course_description(self):
        raw = '劳动教育实践1-生活劳动\n(周边公共设施维护、公共区域管理及景观养护)\n浦金,张秋爽\n12-15(周)[05-06节]\n图书馆'
        courses, errors, warnings = parse_course_cell(raw, 2, 870, 970, 'D6')
        self.assertEqual(errors, [])
        self.assertEqual(warnings, [])
        self.assertIn('周边公共设施', courses[0]['name'])
        self.assertEqual(courses[0]['weeks'], [12, 13, 14, 15])
        for bad in ['坏周次(周)[05-06节]', '12-坏(周)', '12-15(周)[坏节次]']:
            self.assertTrue(parse_course_cell(raw + '\n' + bad, 2, 870, 970, 'D6')[1])

    def test_merge_free_tables(self):
        from core import merge_free_tables, export_template
        with tempfile.TemporaryDirectory() as directory:
            a, b, out = [Path(directory) / name for name in ['a.xlsx', 'b.xlsx', 'out.xlsx']]
            export_template(a, [], [], roster={('海甸', 1, 0, 0): '甲 (1-3)\n乙 (2)'})
            export_template(b, [], [], roster={('海甸', 1, 0, 0): '甲 (3-5)', ('观澜湖', 0, 2, 4): '丙 (1)'})
            merged = merge_free_tables([a, b])
            self.assertEqual(merged[('海甸', 1, 0, 0)], '乙 (2)\n甲 (1-5)')
            export_template(out, [], [], roster=merged)
            self.assertEqual(merge_free_tables([out]), merged)
            export_template(b, [], [], roster={('海甸', 1, 0, 0): '甲'})
            with self.assertRaises(ValueError):
                merge_free_tables([a, b])
            self.assertEqual(merge_free_tables([b, b])[('海甸', 1, 0, 0)], '甲')

    def test_weeks_without_period_marker(self):
        courses, errors, warnings = parse_course_cell('电路理论\n1-4,6-9(周)\n教室', 3, 460, 560, 'E4')
        self.assertEqual(errors, [])
        self.assertTrue(warnings)
        self.assertEqual(courses[0]['weeks'], [1, 2, 3, 4, 6, 7, 8, 9])
        self.assertEqual(courses[0]['periods'], [])
        self.assertEqual(courses[0]['name'], '')
        p = people()[0]
        p.courses, p.warnings = courses, warnings
        self.assertFalse(available(p, 1, 3, 0))
        self.assertTrue(available(p, 5, 3, 0))
        for text in ['课程\n1-4(周)[坏节次]', '课程\n4-1(周)', '课程\n0(周)']:
            self.assertTrue(parse_course_cell(text, 0, 460, 560, 'B4')[1])

    def test_solver_progress_and_cancellation(self):
        import threading
        phases = []
        result, _ = solve_schedule(people(), 1, {'0:0': 1}, [], progress=phases.append)
        self.assertEqual(len(result['0:0']), 1)
        self.assertIn('满足班次人数', phases)
        event = threading.Event()
        def cancel(label):
            event.set()
        with self.assertRaises(InterruptedError):
            solve_schedule(people(), 1, {'0:0': 1}, [], progress=cancel, cancel_event=event)

    def test_multiline_and_time_only_courses(self):
        courses, errors, warnings = parse_course_cell('课程名称\n（分组说明）\n教师甲\n1-5（周）【01-02节】\n地点第一行\n地点第二行', 0, 460, 560, 'B4')
        self.assertEqual(errors, [])
        self.assertEqual(warnings, [])
        self.assertEqual(courses[0]['name'], '课程名称\n（分组说明）')
        self.assertEqual(courses[0]['location'], '地点第一行\n地点第二行')
        courses, errors, warnings = parse_course_cell('1-5(周)[01-02节]', 0, 460, 560, 'B4')
        self.assertEqual(errors, [])
        self.assertTrue(warnings)
        self.assertEqual(courses[0]['name'], '')
        p = people()[0]
        p.courses, p.warnings = courses, warnings
        self.assertFalse(available(p, 1, 0, 0))
        self.assertTrue(available(p, 1, 1, 0))
        result, _ = solve_schedule([p], 1, {'0:0': 1, '1:0': 1}, [])
        self.assertEqual(result, {'0:0': [], '1:0': [p.id]})
        courses, errors, warnings = parse_course_cell('课程\n教师\n1-2(周)[01-02节]\n地点\n另一门\n教师\n3-4(周)[01-02节]\n地点', 0, 460, 560, 'B4')
        self.assertEqual(errors, [])
        self.assertEqual(len(courses), 2)
        self.assertEqual(len(warnings), 2)
        self.assertTrue(all(c['name'] == '' for c in courses))
        _, errors, _ = parse_course_cell('坏周次(周)[01-02节]', 0, 460, 560, 'B4')
        self.assertTrue(errors)

    def test_source(self):
        source = Path(__file__).parent / 'examples' / 'sample_20243006951.xls'
        if not source.exists():
            self.skipTest('Private source timetable is not distributed with the repository')
        p = parse_timetable(str(source))
        self.assertEqual(p.name, '陈锦瀚')
        self.assertEqual(p.id, '20243006951')
        self.assertEqual(len(p.courses), 27)
        self.assertEqual(p.errors, [])
        self.assertEqual(len([c for c in p.courses if c['cell'] == 'F7']), 4)
        self.assertEqual(p.notes, [])
        self.assertTrue(available(p, 1, 5, 0))
        p.notes_confirmed = True
        self.assertTrue(available(p, 1, 5, 0))
        self.assertFalse(available(p, 1, 0, 0))
        self.assertTrue(available(p, 6, 0, 0))
        self.assertFalse(available(p, 12, 0, 0))

    def test_weeks(self):
        self.assertEqual(weeks_from_text('1-3,7,10-12'), [1, 2, 3, 7, 10, 11, 12])
        for text in ['0', '3-1', 'foo', '61']:
            with self.assertRaises(ValueError):
                weeks_from_text(text)

    def test_total_free_and_custom_slots(self):
        p = people()[0]
        p.courses = [dict(day=0, weeks=[10, 11, 13], start=480, end=540)]
        slots = [('自定义', '08:30-10:00')]
        cells = total_free_cells([p], range(1, 19), slots)
        self.assertEqual(cells['0:0'], '甲 (1-9, 12, 14-18)')
        self.assertEqual(compact_weeks([4, 1, 3, 2, 8]), '1-4, 8')
        result, _ = solve_schedule([p], 10, {'0:0': 1}, [], slots)
        self.assertEqual(result['0:0'], [])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'total.xlsx'
            export_grid(path, '总空课表', '1-18', cells, slots)
            book = load_workbook(path)
            self.assertEqual(book.active['B4'].value, '甲 (1-9, 12, 14-18)')
            self.assertEqual(book.active.print_area, "'课表'!$A$1:$H$5")
            book.close()

    def test_exception_confirmation(self):
        p = people()[0]
        p.errors = ['测试异常']
        self.assertFalse(available(p, 1, 0, 0))
        p.errors_confirmed = True
        self.assertTrue(available(p, 1, 0, 0))

    def test_template_preservation(self):
        import zipfile
        import re
        from core import export_template
        p = people()[0]
        p.courses = [dict(day=0, weeks=[10, 11, 13], start=460, end=560)]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'template.xlsx'
            export_template(path, [p], list(range(1, 19)), True)
            book = load_workbook(path)
            self.assertEqual(book['海甸']['C2'].value, '甲 (1-9, 12, 14-18)')
            self.assertEqual(book['海甸']['C2'].font.name, '仿宋_GB2312')
            self.assertEqual(book['海甸']['C3'].value, '')
            book.close()
            with zipfile.ZipFile(Path(__file__).parent / 'resources/free_template.xlsx') as before, zipfile.ZipFile(path) as after:
                for name in before.namelist():
                    a, b = before.read(name), after.read(name)
                    if name in ['xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml']:
                        pattern = rb'<c\b[^>]*\br="[C-I](?:[2-9]|1[0-9]|2[0-6])"[^>]*?(?:/>|>.*?</c>)'
                        a, b = re.sub(pattern, b'', a, flags=re.S), re.sub(pattern, b'', b, flags=re.S)
                    self.assertEqual(a, b, name)

    def test_template_rejects_silent_empty_export(self):
        from core import export_template
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'empty.xlsx'
            with self.assertRaises(ValueError):
                export_template(path, [], [1])
            p = people()[0]
            p.department = '不存在的部门'
            with self.assertRaises(ValueError):
                export_template(path, [p], [1])
            self.assertFalse(path.exists())

    def test_interval_boundaries(self):
        p = people()[0]
        p.courses = [dict(day=0, weeks=[1], start=440, end=460)]
        self.assertTrue(available(p, 1, 0, 0))
        p.courses[0]['end'] = 461
        self.assertFalse(available(p, 1, 0, 0))

    def test_exclusion_and_fixed(self):
        ps = people()
        rules = [rule('exclude', '0'), rule('fixed', '1')]
        result, notes = solve_schedule(ps, 1, {'0:0': 1}, rules)
        self.assertEqual(result['0:0'], ['1'])
        self.assertEqual(notes, [])
        with self.assertRaises(ValueError):
            solve_schedule(ps, 1, {'0:0': 1}, [rule('fixed', '0'), rule('exclude', '0')])
        with self.assertRaises(ValueError):
            solve_schedule(ps, 1, {'0:0': 1}, [rule('fixed', '0'), rule('fixed', '1')])

    def test_preference_priority(self):
        result, _ = solve_schedule(people(), 1, {'0:0': 1}, [rule('prefer', '2'), rule('prefer', '1')])
        self.assertEqual(result['0:0'], ['2'])
        result, _ = solve_schedule(people(), 1, {'0:0': 1}, [rule('prefer', '1'), rule('prefer', '2')])
        self.assertEqual(result['0:0'], ['1'])

    def test_multi_person_rules(self):
        r = dict(kind='fixed', people=['0', '1'], day=0, slot=0, weeks=[], enabled=True)
        result, _ = solve_schedule(people(), 1, {'0:0': 2}, [r])
        self.assertEqual(set(result['0:0']), {'0', '1'})
        with self.assertRaises(ValueError):
            solve_schedule(people(), 1, {'0:0': 1}, [r])
        r['kind'] = 'exclude'
        self.assertEqual(candidates(people(), 1, 0, 0, [r]), ['2'])
        r['kind'] = 'prefer'
        result, _ = solve_schedule(people(), 1, {'0:0': 2}, [r])
        self.assertEqual(set(result['0:0']), {'0', '1'})

    def test_balance_and_shortage(self):
        result, _ = solve_schedule(people(), 1, {'0:0': 1, '1:0': 1, '2:0': 1}, [])
        self.assertEqual(sorted(pid for ids in result.values() for pid in ids), ['0', '1', '2'])
        result, _ = solve_schedule(people()[:1], 1, {'0:0': 2}, [])
        self.assertEqual(result['0:0'], ['0'])

    def test_overlap_and_fixed_conflict(self):
        slots = list(SLOTS)
        slots[1] = ('重叠时段', '08:00-10:00')
        result, _ = solve_schedule(people()[:1], 1, {'0:0': 1, '0:1': 1}, [], slots)
        self.assertEqual(sum(len(ids) for ids in result.values()), 1)
        with self.assertRaises(ValueError):
            solve_schedule(people()[:1], 1, {'0:0': 1, '0:1': 1}, [rule('fixed', '0', slot=0), rule('fixed', '0', slot=1)], slots)

    def test_week_specific_rules(self):
        ps = people()
        r = rule('exclude', '0', weeks=[2])
        self.assertIn('0', candidates(ps, 1, 0, 0, [r]))
        self.assertNotIn('0', candidates(ps, 2, 0, 0, [r]))

    def test_export_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'table.xlsx'
            export_grid(path, '空课表', 5, {'0:0': '甲\n乙'})
            b = load_workbook(path)
            s = b.active
            self.assertEqual(s['B4'].value, '甲\n乙')
            self.assertEqual(s['A8'].value, '9、10、11节\n19:20-21:55')
            self.assertIn('A1:H1', [str(r) for r in s.merged_cells.ranges])
            self.assertEqual(s.freeze_panes, 'B4')
            b.close()


if __name__ == '__main__':
    unittest.main()
